#!/usr/bin/env python3
"""Validate final reconciliation details and update the Bank_Statement sheet."""

from __future__ import annotations

import argparse
import sys
from copy import copy
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Color, Font, PatternFill
from openpyxl.utils import get_column_letter

from reconcile import (
    GREEN_FILL,
    clean_string,
    find_last_bank_data_row,
    locate_bank_header,
    money_equal,
    numeric_amount,
    parse_date,
    reference_text,
)

DETAILS_NOT_FOUND = "Details not found in reconciliation sheets"
RED_FILL = PatternFill(fill_type="solid", fgColor=Color(rgb="FFFFC7CE"))
YELLOW_FILL = PatternFill(fill_type="solid", fgColor=Color(rgb="FFFFEB9C"))
AMOUNT_FORMAT = "#,##0.00"


def reference_key(value: Any) -> str:
    """Normalize reference IDs for comparison while retaining their displayed text."""
    reference = reference_text(value).upper()
    return reference.lstrip("0") or "0" if reference.isdigit() else reference


def header_columns(ws, header_row: int = 1) -> dict[str, int]:
    return {
        clean_string(ws.cell(header_row, col).value).upper(): col
        for col in range(1, ws.max_column + 1)
        if clean_string(ws.cell(header_row, col).value)
    }


def add_match(matches: dict[str, list[dict[str, Any]]], reference: Any, amount: float,
              direction: str, description: str, source: str, source_row: int | None = None,
              transaction_date: Any = None) -> None:
    normalized_reference = reference_key(reference)
    if not normalized_reference:
        return
    matches.setdefault(normalized_reference, []).append({
        "amount": amount,
        "direction": direction,
        "description": description,
        "source": source,
        "source_row": source_row,
        "transaction_date": parse_date(transaction_date),
    })


def expense_payment_amount(ws, row: int) -> float:
    taxable = numeric_amount(ws.cell(row, 6).value)
    sgst = numeric_amount(ws.cell(row, 7).value)
    cgst = numeric_amount(ws.cell(row, 8).value)
    tds_rate = numeric_amount(ws.cell(row, 10).value)
    calculated_payment = round(taxable + sgst + cgst - round(taxable * tds_rate / 100, 0), 0)
    amount_paid_value = ws.cell(row, 12).value
    amount_paid = numeric_amount(amount_paid_value)
    if amount_paid > 0 and not (isinstance(amount_paid_value, str) and amount_paid_value.startswith("=")):
        return amount_paid
    return calculated_payment


def collect_expense_matches(ws, matches: dict[str, list[dict[str, Any]]]) -> None:
    for row in range(2, ws.max_row + 1):
        vendor = clean_string(ws.cell(row, 5).value)
        if not vendor:
            break
        bill_number = clean_string(ws.cell(row, 4).value)
        add_match(
            matches,
            ws.cell(row, 16).value,
            expense_payment_amount(ws, row),
            "debit",
            f"Expense | {bill_number} | {vendor}",
            "Expense",
            row,
        )


def collect_income_matches(ws, matches: dict[str, list[dict[str, Any]]]) -> None:
    columns = header_columns(ws)
    reference_col = columns.get("BANK REFERENCE ID")
    amount_col = columns.get("CREDIT")
    flat_col = columns.get("FLAT NO.")
    if not reference_col or not amount_col:
        return
    for row in range(2, ws.max_row + 1):
        amount = numeric_amount(ws.cell(row, amount_col).value)
        if amount <= 0:
            continue
        flat = clean_string(ws.cell(row, flat_col).value) if flat_col else ""
        add_match(matches, ws.cell(row, reference_col).value, amount, "credit",
                  f"Income recorded in NBH | {flat}".rstrip(" |"), "Income_recorded_in_nbh")


def collect_other_income_matches(ws, matches: dict[str, list[dict[str, Any]]]) -> None:
    for row in range(2, ws.max_row + 1):
        reference_cell = ws.cell(row, 14)
        reference = reference_text(reference_cell.value)
        if reference:
            reference_cell.value = reference
            reference_cell.number_format = "@"
        amount = numeric_amount(ws.cell(row, 13).value)
        if amount <= 0:
            amount = numeric_amount(ws.cell(row, 5).value)
        if not reference or amount <= 0:
            continue
        income_type = clean_string(ws.cell(row, 12).value)
        if not income_type:
            continue
        ledger_name = clean_string(ws.cell(row, 2).value)
        description = f"Income from Other Sources | {income_type}"
        if ledger_name:
            description += f" | {ledger_name}"
        transaction_date = ws.cell(row, 15).value or ws.cell(row, 3).value
        add_match(matches, reference, amount, "credit", description,
                  "Income from Other Sources", row, transaction_date)


def collect_fd_matches(ws, matches: dict[str, list[dict[str, Any]]]) -> None:
    for row in range(2, ws.max_row + 1):
        reference = ws.cell(row, 3).value
        debit = numeric_amount(ws.cell(row, 6).value)
        credit = numeric_amount(ws.cell(row, 7).value)
        transaction_type = clean_string(ws.cell(row, 2).value)
        if debit > 0:
            add_match(matches, reference, debit, "debit", f"FD Transactions | {transaction_type}", "FD Transactions")
        if credit > 0:
            add_match(matches, reference, credit, "credit", f"FD Transactions | {transaction_type}", "FD Transactions")


def collect_fd_interest_matches(ws, matches: dict[str, list[dict[str, Any]]]) -> None:
    for row in range(2, ws.max_row):
        amount = numeric_amount(ws.cell(row, 5).value)
        if amount > 0:
            add_match(matches, ws.cell(row, 4).value, amount, "credit",
                      "Income from FD Interest", "Income from FD Interest")


def collect_cash_matches(workbook, matches: dict[str, list[dict[str, Any]]]) -> None:
    for sheet_name in workbook.sheetnames:
        if not sheet_name.startswith("Cash Withdrawal "):
            continue
        ws = workbook[sheet_name]
        reference = ws["E1"].value
        withdrawal_amount = numeric_amount(ws["H1"].value)
        opening_balance = numeric_amount(ws["B3"].value)
        payment_header_row = next(
            (row for row in range(1, ws.max_row + 1)
             if any(clean_string(ws.cell(row, col).value).upper() == "PAYMENT AMOUNT"
                    for col in range(1, ws.max_column + 1))),
            None,
        )
        total_paid_row = next(
            (row for row in range(1, ws.max_row + 1)
             if any(clean_string(ws.cell(row, col).value).upper() == "TOTAL PAID"
                    for col in range(1, ws.max_column + 1))),
            None,
        )
        if not payment_header_row or not total_paid_row:
            continue
        payment_col = next(
            col for col in range(1, ws.max_column + 1)
            if clean_string(ws.cell(payment_header_row, col).value).upper() == "PAYMENT AMOUNT"
        )
        payment_start_row = payment_header_row + 1
        payment_total = sum(
            numeric_amount(ws.cell(row, payment_col).value)
            for row in range(payment_start_row, total_paid_row)
        )
        payment_entries = sum(
            1 for row in range(payment_start_row, total_paid_row)
            if numeric_amount(ws.cell(row, payment_col).value) > 0
        )
        closing_label_col = next(
            (col for col in range(1, ws.max_column + 1)
             if clean_string(ws.cell(total_paid_row, col).value).upper() == "CLOSING BALANCE"),
            None,
        )
        closing_value = ws.cell(total_paid_row, closing_label_col + 1).value if closing_label_col else None
        closing_balance = numeric_amount(closing_value)
        if isinstance(closing_value, str) and closing_value.startswith("="):
            closing_balance = opening_balance + withdrawal_amount - payment_total
        accounted_total = payment_total + closing_balance - opening_balance
        if withdrawal_amount > 0 and payment_entries > 0:
            fully_accounted = money_equal(accounted_total, withdrawal_amount, tolerance=1.0)
            add_match(matches, reference, withdrawal_amount, "debit",
                      f"Cash withdrawal | {sheet_name}", sheet_name)
            matches[reference_key(reference)][-1]["cash_state"] = "MATCHED" if fully_accounted else "PARTIAL"
            matches[reference_key(reference)][-1]["accounted_total"] = accounted_total


def fill_bank_row(bank_ws, row: int, fill: PatternFill) -> None:
    for col in range(1, bank_ws.max_column + 1):
        bank_ws.cell(row, col).fill = copy(fill)


def add_bank_totals(bank_ws, header_row: int, last_data_row: int, debit_col: int,
                    credit_col: int, description_col: int) -> None:
    total_row = last_data_row + 2
    bank_ws.cell(total_row, description_col).value = "Total"
    debit_letter = bank_ws.cell(1, debit_col).column_letter
    credit_letter = bank_ws.cell(1, credit_col).column_letter
    bank_ws.cell(total_row, debit_col).value = f"=SUM({debit_letter}{header_row + 1}:{debit_letter}{last_data_row})"
    bank_ws.cell(total_row, credit_col).value = f"=SUM({credit_letter}{header_row + 1}:{credit_letter}{last_data_row})"
    for col in (debit_col, credit_col, description_col):
        cell = bank_ws.cell(total_row, col)
        cell.font = Font(bold=True)
        if col != description_col:
            cell.number_format = AMOUNT_FORMAT


def autofit_workbook_columns(workbook, max_width: int = 80) -> None:
    """Apply the shared compact widths and wrapped multiline formatting."""
    from reconcile import autofit

    for ws in workbook.worksheets:
        autofit(ws)


def build_match_index(workbook) -> dict[str, list[dict[str, Any]]]:
    matches: dict[str, list[dict[str, Any]]] = {}
    collectors = (
        ("Expense", collect_expense_matches),
        ("Income_recorded_in_nbh", collect_income_matches),
        ("Income from Other Sources", collect_other_income_matches),
        ("FD Transactions", collect_fd_matches),
        ("Income from FD Interest", collect_fd_interest_matches),
    )
    for sheet_name, collector in collectors:
        if sheet_name in workbook.sheetnames:
            collector(workbook[sheet_name], matches)
    collect_cash_matches(workbook, matches)
    return matches


def update_detail_match_status(workbook, item: dict[str, Any]) -> None:
    """Update the originating detail row only for Expense and Other Sources."""
    source = item["source"]
    source_row = item["source_row"]
    if source not in {"Expense", "Income from Other Sources"} or source_row is None:
        return

    ws = workbook[source]
    columns = header_columns(ws)
    status_col = columns.get("BANK MATCH STATUS") or columns.get("MATCH STATUS")
    if not status_col:
        status_col = ws.max_column + 1
        ws.cell(1, status_col).value = "Bank Match Status"
    ws.cell(source_row, status_col).value = "MATCHED"
    for col in range(1, ws.max_column + 1):
        ws.cell(source_row, col).fill = copy(GREEN_FILL)


def apply_final_reconciliation(workbook) -> tuple[int, int, int]:
    if "Bank_Statement" not in workbook.sheetnames:
        raise ValueError("Workbook does not contain a Bank_Statement sheet.")

    bank_ws = workbook["Bank_Statement"]
    header_row, headers = locate_bank_header(bank_ws)
    date_col = headers.get("TRANSACTION DATE", 1)
    narration_col = headers.get("NARRATION", 2)
    debit_col = headers.get("DEBIT", 4)
    credit_col = headers.get("CREDIT", 5)
    reference_col = headers.get("GENERATED REFERENCE ID", 6)
    description_col = headers.get("EXPENSE DESCRIPTION", 8)
    status_col = headers.get("MATCH STATUS", 9)
    last_row = find_last_bank_data_row(bank_ws, header_row)
    match_index = build_match_index(workbook)
    other_income_candidates = [
        item for items in match_index.values() for item in items
        if item["source"] == "Income from Other Sources"
    ]

    matched = 0
    details_missing = 0
    amount_mismatch = 0
    for row in range(header_row + 1, last_row + 1):
        reference = reference_key(bank_ws.cell(row, reference_col).value)
        debit = numeric_amount(bank_ws.cell(row, debit_col).value)
        credit = numeric_amount(bank_ws.cell(row, credit_col).value)
        direction = "debit" if debit > 0 else "credit"
        bank_amount = debit if debit > 0 else credit
        existing_description = clean_string(bank_ws.cell(row, description_col).value)
        existing_status = clean_string(bank_ws.cell(row, status_col).value).upper()

        # Keep transaction details that were already validated in the draft workbook.
        if existing_status == "MATCHED" and existing_description and existing_description != DETAILS_NOT_FOUND:
            existing_candidates = [
                item for item in match_index.get(reference, [])
                if item["direction"] == direction and money_equal(item["amount"], bank_amount, tolerance=1.0)
            ]
            if len(existing_candidates) == 1:
                update_detail_match_status(workbook, existing_candidates[0])
            for col in range(1, bank_ws.max_column + 1):
                bank_ws.cell(row, col).fill = copy(GREEN_FILL)
            matched += 1
            continue

        candidates = [
            item for item in match_index.get(reference, [])
            if item["direction"] == direction and money_equal(item["amount"], bank_amount, tolerance=1.0)
        ]
        if not candidates:
            bank_date = parse_date(bank_ws.cell(row, date_col).value)
            candidates = [
                item for item in other_income_candidates
                if item["direction"] == direction
                and item["transaction_date"] == bank_date
                and money_equal(item["amount"], bank_amount, tolerance=1.0)
            ]

        if len(candidates) == 1:
            item = candidates[0]
            if item.get("cash_state") == "PARTIAL":
                bank_ws.cell(row, description_col).value = (
                    f'{item["description"]} | Account partially matched: '
                    f'accounted {item["accounted_total"]:,.2f} of {bank_amount:,.2f}'
                )
                bank_ws.cell(row, status_col).value = "PARTIALLY MATCHED"
                fill_bank_row(bank_ws, row, YELLOW_FILL)
                amount_mismatch += 1
                continue
            bank_ws.cell(row, description_col).value = item["description"]
            bank_ws.cell(row, status_col).value = "MATCHED"
            update_detail_match_status(workbook, item)
            fill_bank_row(bank_ws, row, GREEN_FILL)
            matched += 1
        elif reference in match_index:
            bank_ws.cell(row, description_col).value = "Details found, but amount or transaction type does not match"
            bank_ws.cell(row, status_col).value = "AMOUNT MISMATCH"
            fill_bank_row(bank_ws, row, RED_FILL)
            amount_mismatch += 1
        else:
            bank_ws.cell(row, description_col).value = DETAILS_NOT_FOUND
            bank_ws.cell(row, status_col).value = "DETAILS NOT FOUND"
            fill_bank_row(bank_ws, row, RED_FILL)
            details_missing += 1

    add_bank_totals(bank_ws, header_row, last_row, debit_col, credit_col, description_col)
    return matched, details_missing, amount_mismatch


def main() -> None:
    parser = argparse.ArgumentParser(description="Validate and finalize an NBH reconciliation workbook.")
    parser.add_argument("--input", required=True, help="Output workbook created by reconcile.py")
    parser.add_argument("--output", help="Final reconciled workbook path (.xlsx)")
    args = parser.parse_args()

    input_path = Path(args.input)
    if not input_path.exists() or input_path.suffix.lower() != ".xlsx":
        print(f"ERROR: Input workbook not found or not an .xlsx file: {input_path}")
        sys.exit(1)
    output_path = Path(args.output) if args.output else input_path.with_name(f"{input_path.stem}_FINAL.xlsx")

    workbook = load_workbook(input_path, data_only=False)
    try:
        matched, details_missing, amount_mismatch = apply_final_reconciliation(workbook)
        autofit_workbook_columns(workbook)
        workbook.save(output_path)
    finally:
        workbook.close()

    print(f"Saved final reconciliation: {output_path}")
    print(f"Transactions matched       : {matched}")
    print(f"Details not found          : {details_missing}")
    print(f"Amount/type mismatch       : {amount_mismatch}")


if __name__ == "__main__":
    main()
