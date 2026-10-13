"""FY27 (2026-27) dashboard: reads the monthly report_*.xlsx files produced by auto_reconcile.py."""
import datetime as dt
import glob
import os
import re

import openpyxl

from dashboard_common import build_page, categorize_expense, fmt_date, num

MONTHS = ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"]
LIFT_AMC_MONTHLY = 904999.0 / 12


def is_matched(status):
    text = str(status or "").upper()
    # FD transfers are reconciled through the FD Transactions sheet.
    return text.startswith("MATCHED") or text.startswith("FD TRANSFER")


def status_bucket(status):
    text = str(status or "").strip().upper()
    if is_matched(text):
        return "Matched"
    if text.startswith("PARTIALLY"):
        return "Partially matched"
    if text.startswith("UNMATCHED"):
        return "Unmatched"
    if text.startswith("PENDING"):
        return "Pending"
    if text.startswith("NOT EXPENSE"):
        return "Not expense / no income entry"
    return "Review"


def row_class(status):
    return "warn" if status_bucket(status) == "Partially matched" else "bad"


def count_buckets(statuses):
    counts = {}
    for status in statuses:
        label = status_bucket(status)
        counts[label] = counts.get(label, 0) + 1
    return [{"label": k, "value": v} for k, v in counts.items()]


def month_sort_key(path):
    name = os.path.basename(os.path.dirname(path))
    match = re.match(r"([A-Za-z]{3})", name)
    year = re.search(r"(\d{4}|\d{2})$", name)
    month = MONTHS.index(match.group(1).lower()) + 1 if match and match.group(1).lower() in MONTHS else 0
    y = int(year.group(1)) if year else 0
    y = y + 2000 if y < 100 else y
    return (y, month)


def month_label(path):
    y, m = month_sort_key(path)
    if m:
        return dt.date(y or 2000, m, 1).strftime("%b %Y") if y else dt.date(2000, m, 1).strftime("%b")
    return os.path.basename(os.path.dirname(path))


def rows_until(ws, stop):
    for row in ws.iter_rows(min_row=2, values_only=True):
        if stop(row):
            break
        yield row


def read_bank(ws):
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if isinstance(row[0], (dt.datetime, dt.date)) and row[8] is not None:
            rows.append(row)
    return rows


def read_income_other(ws):
    out = []
    for row in rows_until(ws, lambda r: r[0] == "Income Type"):
        if row[4] is None or row[11] is None:
            continue
        out.append(row)
    return out


def income_base(row):
    """Amount without GST, using the same formula as the reconciliation report."""
    taxable = row[6]
    if taxable not in (None, ""):
        return num(taxable)
    pct = num(row[7])
    return round((num(row[4]) - num(row[5])) / (1 + 2 * pct), 2)


def read_expenses(ws):
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] == "Billing Category":
            break
        if row[4] is None or row[17] is None:
            continue
        out.append(row)
    return out


def expense_total(row):
    invoice = round(num(row[5]) + num(row[6]) + num(row[7]))
    if invoice:
        return float(invoice)
    return num(row[19]) or num(row[11])


def normalized_bill_number(value):
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def display_text(value):
    text = re.sub(r"\s*RS\s*\(\s*3780\s+PAID\s+BY\s+KISHORE\s*\)\s*", " ", str(value or ""), flags=re.I)
    return " ".join(text.split())


def report_expense_category(name):
    text = str(name or "").lower()
    if "gowthami" in text and any(term in text for term in ("stp system service", "sand filter service")):
        return "Others"
    return categorize_expense(name)


def bank_debits(bank, report_month):
    gross = fd = reimbursement = 0.0
    vendor_exclusions = {}
    reference_exclusions = {
        "Apr 2026": {
            "807960": "Gowthami sand filter service",
            "807961": "Gowthami bill #945",
        },
    }.get(report_month, {})
    for row in bank:
        amount = num(row[3])
        if amount <= 0:
            continue
        narration = str(row[1] or "").upper()
        description = str(row[7] or "").upper()
        status = str(row[8] or "").upper()
        detail = " ".join((narration, description, status))
        gross += amount
        if description.startswith("FD ") or "TO FD" in narration:
            fd += amount
        elif description.startswith("REIMBURSEMENT") or "REIMBURS" in narration:
            reimbursement += amount
        elif str(row[5] or "").strip() in reference_exclusions:
            label = reference_exclusions[str(row[5]).strip()]
            vendor_exclusions[label] = vendor_exclusions.get(label, 0.0) + amount
        elif "OTIS" in detail:
            vendor_exclusions["OTIS"] = vendor_exclusions.get("OTIS", 0.0) + amount
        elif "IT DEPARTMENT" in detail:
            vendor_exclusions["IT Department"] = vendor_exclusions.get("IT Department", 0.0) + amount
    exclusions = []
    if fd:
        exclusions.append({"label": "FD transactions", "amount": round(fd, 2)})
    if reimbursement:
        exclusions.append({"label": "Sponsorship payout", "amount": round(reimbursement, 2)})
    exclusions.extend(
        {"label": label, "amount": round(amount, 2)}
        for label, amount in vendor_exclusions.items()
    )
    excluded_total = fd + reimbursement + sum(vendor_exclusions.values())
    return {
        "gross": round(gross, 2),
        "fd": round(fd, 2),
        "reimbursement": round(reimbursement, 2),
        "exclusions": exclusions,
        "excluded_total": round(excluded_total, 2),
        "net": round(gross - excluded_total, 2),
    }


def process_report(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    lift_amc = LIFT_AMC_MONTHLY if os.path.basename(os.path.dirname(os.path.dirname(path))) == "2026-2027" else 0.0
    bank = read_bank(wb["Bank_Statement"])
    nbh = [
        r for r in wb["Income_recorded_in_nbh"].iter_rows(min_row=2, values_only=True)
        if r[1] is not None and r[13] is not None
    ]
    other = read_income_other(wb["Income from Other Sources"])
    expenses = read_expenses(wb["Expense"])

    income_cat = {}
    for row in other:
        label = str(row[11])
        base = income_base(row)
        if label.lower() == "sponsorship":
            base = round(base * 0.10, 2)
        income_cat[label] = income_cat.get(label, 0.0) + base
    income_cat = {k: round(v, 2) for k, v in income_cat.items() if v}
    other_income = round(sum(income_cat.values()), 2)

    fd_rows = [
        r for r in wb["FD Transactions"].iter_rows(min_row=2, values_only=True)
        if r[1] == "FD Interest" and num(r[6]) > 0
    ]
    fd_income = round(sum(num(r[6]) for r in fd_rows), 2)

    # Vendor bills are gross; tax payments are excluded to avoid double counting.
    expense_cat = {}
    others_breakup = {}
    paid_refs = set()
    april_excluded_bills = {"35", "945", "944", "inv00087"}
    report_month = month_label(path)
    excluded_vendors = {
        "Jun 2026": ("otis",),
        "Jul 2026": ("it department",),
    }
    for row in expenses:
        for ref in re.split(r"[,\s]+", str(row[15] or "")):
            if ref:
                paid_refs.add(ref)
        if row[18] == "Tax and Duties":
            continue
        vendor_name = str(row[4] or "").lower()
        if report_month == "Apr 2026" and normalized_bill_number(row[3]) in april_excluded_bills:
            continue
        if any(name in vendor_name for name in excluded_vendors.get(report_month, ())):
            continue
        label = report_expense_category(row[4])
        expense_cat[label] = expense_cat.get(label, 0.0) + expense_total(row)
        if label == "Others":
            vendor = display_text(row[4])
            others_breakup[vendor] = others_breakup.get(vendor, 0.0) + expense_total(row)

    # Cash withdrawals not already represented by an expense row
    cash_total = 0.0
    for row in bank:
        narration = str(row[1] or "").upper()
        if num(row[3]) > 0 and ("CASH WITHDRAWAL" in narration or "CASH CHEQUE" in narration):
            if str(row[5] or "") in paid_refs:
                continue
            # Matched rows are skipped unless the report itself labels them "Cash withdrawal".
            if is_matched(row[8]) and not str(row[7] or "").lower().startswith("cash withdrawal"):
                continue
            cash_total += num(row[3])
    if cash_total:
        expense_cat["Cash Withdrawals"] = expense_cat.get("Cash Withdrawals", 0.0) + cash_total
    expense_cat = {k: round(v, 2) for k, v in expense_cat.items() if v}
    total_expense = round(sum(expense_cat.values()), 2)

    unreconciled_bank = [
        {
            "date": fmt_date(r[0]),
            "narration": display_text(r[1]),
            "ref": str(r[5] or ""),
            "debit": num(r[3]),
            "credit": num(r[4]),
            "status": str(r[8]),
            "cls": row_class(r[8]),
        }
        for r in bank if not is_matched(r[8])
    ]
    unreconciled_expense = [
        {
            "bill": str(row[3] or "").strip(),
            "name": display_text(row[4]),
            "billing_month": str(row[2] or ""),
            "invoice_total": expense_total(row),
            "bank_paid": num(row[19]),
            "status": str(row[17]),
            "cls": row_class(row[17]),
        }
        for row in expenses
        if not is_matched(row[17]) and row[18] != "Tax and Duties"
    ]

    return {
        "label": month_label(path),
        "debits": bank_debits(bank, month_label(path)),
        "bank_pie": count_buckets(r[8] for r in bank),
        "income_pie": count_buckets([r[13] for r in nbh] + [r[15] for r in other]),
        "income_categories": [{"label": k, "value": v} for k, v in sorted(income_cat.items(), key=lambda x: -x[1])],
        "fd_income": fd_income,
        "fd_entries": [{"date": fmt_date(r[0]), "ref": str(r[2] or ""), "amount": num(r[6])} for r in fd_rows],
        "expense_categories": [{"label": k, "value": v} for k, v in sorted(expense_cat.items(), key=lambda x: -x[1])],
        "others_breakup": [
            {"label": k, "value": round(v, 2)}
            for k, v in sorted(others_breakup.items(), key=lambda x: -x[1]) if v
        ],
        "summary": {
            "total_expense": total_expense,
            "lift_amc": round(lift_amc, 2),
            "total_with_lift_amc": round(total_expense + lift_amc, 2),
            "recurring_expense": round(sum(v for k, v in expense_cat.items() if k not in ("Others", "Cash Withdrawals")), 2),
            "other_expense": round(expense_cat.get("Others", 0.0), 2),
            "cash_withdrawals": round(expense_cat.get("Cash Withdrawals", 0.0), 2),
        },
        "unreconciled_bank": unreconciled_bank,
        "unreconciled_expense": unreconciled_expense,
    }


BODY_JS = """
function summary(s){
  const row=(l,v,cls)=>`<tr class="${cls||''}"><td>${l}</td><td class="n">${inr(v)}</td><td class="n">${(v/SQFT).toFixed(4)}</td></tr>`;
  return `<table><tr><th>Item</th><th class="n">Amount (&#8377;)</th><th class="n">Per sqft (&#8377;, ${SQFT.toLocaleString("en-IN")} sqft)</th></tr>`
        +row("Recurring",s.recurring_expense)+row("Other",s.other_expense)+row("Cash Withdrawals",s.cash_withdrawals)
    +row("Total expense",s.total_expense,"total")+row("Lift AMC",s.lift_amc)
    +row("Total including lift AMC",s.total_with_lift_amc,"total")+'</table>';
}
function panel(m){
  const cnt=v=>v+"";
  return '<div class="grid">'
   +card("1. Bank statement reconciliation (transactions)",pie(m.bank_pie,(i)=>STATUS_COLORS[i.label],cnt))
   +card("2. Income sheet reconciliation (entries)",pie(m.income_pie,(i)=>STATUS_COLORS[i.label],cnt))
   +card("3. Income by category (without GST; Sponsorship at 10%)",bars(m.income_categories))
   +card("4. Income from FD interest",`<p><b>&#8377; ${inr(m.fd_income)}</b> from ${m.fd_entries.length} entries</p>`+table([{h:"Date",k:"date"},{h:"FD reference",k:"ref"},{h:"Interest (&#8377;)",k:"amount",n:1}],m.fd_entries,"No FD interest this month"))
   +card("5. Expense per category",bars(m.expense_categories))
   +card("6. Breakup of Others expense",bars(m.others_breakup))
   +card("7. Summary and per sqft rate",summary(m.summary))
   +card("8. Bank transactions not yet reconciled",table([{h:"Date",k:"date"},{h:"Narration",k:"narration"},{h:"Reference",k:"ref"},{h:"Debit",k:"debit",n:1},{h:"Credit",k:"credit",n:1},{h:"Status",k:"status"}],m.unreconciled_bank,"All bank transactions reconciled"),1)
   +card("9. Expenses not yet fully reconciled",table([{h:"Bill / Invoice",k:"bill"},{h:"Vendor",k:"name"},{h:"Billing month",k:"billing_month"},{h:"Invoice total",k:"invoice_total",n:1},{h:"Bank paid",k:"bank_paid",n:1},{h:"Status",k:"status"}],m.unreconciled_expense,"All expenses reconciled"),1)
   +'</div>';
}
function renderYear(){
  DATA.forEach(m=>addTab(m.label,panel(m)));
  addTab("Monthly Debits",debitsPanel());
    show(0);
}
"""


def load_year(fy_root):
    reports = sorted(glob.glob(os.path.join(fy_root, "*", "report_*.xlsx")), key=month_sort_key)
    months = []
    for path in reports:
        print("Reading", path)
        months.append(process_report(path))
    return months


def write_dashboard(datasets, output):
    with open(output, "w", encoding="utf-8") as handle:
        handle.write(build_page(datasets, BODY_JS))
    print("Wrote", output)

