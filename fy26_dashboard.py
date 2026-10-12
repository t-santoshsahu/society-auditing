#!/usr/bin/env python3
"""Standalone FY26 (Apr 2025 - Mar 2026) dashboard, kept separate from the FY27+ generator.

Usage:
  python fy26_dashboard.py [--workbook FY26_Expense_Bank_Reconciliation.xlsx]
                           [--petty PettyCash.xlsx] [-o dashboard_fy26.html]
"""
import argparse
import os
import re

import openpyxl

from dashboard_common import build_page, categorize_expense, fmt_date, num

MONTH_NAMES = ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"]
LABEL = "FY26"


def month_key(text):
    """'Apr 2025' -> (2025, 4); PettyCash sheet names such as 'Aug-25' or 'Sept-25' work too."""
    match = re.match(r"^\s*([A-Za-z]{3})[A-Za-z]*[\s-]*(\d{2}|\d{4})\s*$", str(text))
    if not match or match.group(1).lower() not in MONTH_NAMES:
        return None
    year = int(match.group(2))
    return (year + 2000 if year < 100 else year, MONTH_NAMES.index(match.group(1).lower()) + 1)


def petty_cash_received(path):
    """{(year, month): amount received} for months whose petty cash sheet has a 'Received' amount."""
    if not path or not os.path.exists(path):
        return None
    received = {}
    wb = openpyxl.load_workbook(path, data_only=True)
    for ws in wb:
        key = month_key(ws.title)
        if not key:
            continue
        for row in ws.iter_rows(min_row=2, values_only=True):
            if re.match(r"\s*rec(ei)?v", str(row[3] or ""), re.I) and num(row[4]) > 0:
                received[key] = num(row[4])
                break
    return received


def allocate_all(month_order, cash, petty):
    """Match petty cash received against withdrawals, month by month.

    Received money can explain an exact-amount withdrawal from an earlier month that is still unexplained
    (cash drawn late in a month is often recorded in the next month's petty sheet).
    """
    carry = []
    for month in month_order:
        rows = cash.get(month, [])
        for w in rows:
            w.update(covered=0.0, via=None, month=month)
        remaining = (petty or {}).get(month_key(month)) or 0.0
        # Own-month withdrawals are tried first; only then are earlier unexplained ones.
        for w in rows + carry:
            if remaining > 0 and w["covered"] == 0 and abs(w["amount"] - remaining) < 1:
                w["covered"], w["via"] = w["amount"], month
                remaining -= w["amount"]
        for w in rows:
            if w["covered"] == 0 and w["amount"] <= remaining + 1:
                w["covered"], w["via"] = w["amount"], month
                remaining -= w["amount"]
        for w in rows:
            if w["covered"] == 0 and remaining > 0:
                w["covered"], w["via"] = min(remaining, w["amount"]), month
                remaining -= w["covered"]
        carry = [w for w in carry + rows if w["covered"] == 0]

    for rows in cash.values():
        for w in rows:
            full = w["covered"] >= w["amount"] - 1
            if full:
                w["status"] = "Details available" if w["via"] == w["month"] else "Details in %s petty sheet" % w["via"]
            elif w["covered"]:
                w["status"] = "Details for %s of %s" % (format(w["covered"], ",.0f"), format(w["amount"], ",.0f"))
            else:
                w["status"] = "No details found"
            w["cls"] = "ok" if full else "bad"


def process_fy26(workbook_path, petty_path):
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    month_order = [str(r[0]) for r in wb["Monthly Summary"].iter_rows(min_row=2, values_only=True) if r[0]]
    petty = petty_cash_received(petty_path)

    bills, others, cash, debit_rows = {}, {}, {}, {}
    for row in wb["Expense Matching"].iter_rows(min_row=2, values_only=True):
        if not row[0] or not row[3]:
            continue
        # TDS/GST remittance rows are tax payments, not expenses.
        if re.match(r"\s*(paid\s+)?(tds|gst)\b|.*your\s*self.*gst", str(row[3]), re.I):
            continue
        amount = num(row[7]) or num(row[6])
        label = categorize_expense(row[3])
        cats = bills.setdefault(str(row[0]), {})
        cats[label] = cats.get(label, 0.0) + amount
        if label == "Others":
            vendor = " ".join(str(row[3]).split())
            vendors = others.setdefault(str(row[0]), {})
            vendors[vendor] = vendors.get(vendor, 0.0) + amount
    for row in wb["Cash & Unmatched"].iter_rows(min_row=2, values_only=True):
        if row[0] and str(row[1]).lower() == "bank cash withdrawal":
            cash.setdefault(str(row[0]), []).append(
                {"date": fmt_date(row[2]), "description": " ".join(str(row[3] or "").split()), "amount": num(row[4])}
            )
    for row in wb["Bank Debits"].iter_rows(min_row=2, values_only=True):
        if row[1] and num(row[3]) > 0:
            debit_rows.setdefault(str(row[1]), []).append(row)

    months = []
    allocate_all(month_order, cash, petty)
    for month in month_order:
        cats = dict(bills.get(month, {}))
        bill_total = sum(cats.values())
        withdrawals = cash.get(month, [])
        cash_total = sum(w["amount"] for w in withdrawals)
        if cash_total:
            cats["Cash Withdrawals"] = cash_total

        received = None if petty is None else petty.get(month_key(month))
        missing = round(sum(w["amount"] - w["covered"] for w in withdrawals), 2)
        if not withdrawals:
            status = "No cash withdrawals"
        elif missing:
            status = "No details for %s of %s withdrawn" % (format(missing, ",.0f"), format(cash_total, ",.0f"))
        else:
            status = "Details available"
        flagged = missing > 0
        gross = fd = 0.0
        for row in debit_rows.get(month, []):
            gross += num(row[3])
            if "TO FD" in " ".join(str(row[2] or "").upper().split()):
                fd += num(row[3])
        months.append({
            "label": month,
            "expense_categories": [
                {"label": k, "value": round(v, 2)} for k, v in sorted(cats.items(), key=lambda x: -x[1]) if v
            ],
            "others_breakup": [
                {"label": k, "value": round(v, 2)}
                for k, v in sorted(others.get(month, {}).items(), key=lambda x: -x[1]) if v
            ],
            "summary": {
                "bill_expense": round(bill_total, 2),
                "cash": round(cash_total, 2),
                "total_expense": round(bill_total + cash_total, 2),
            },
            "debits": {"gross": round(gross, 2), "fd": round(fd, 2), "reimbursement": 0.0, "net": round(gross - fd, 2)},
            "cash_details": {
                "withdrawals": withdrawals,
                "total": round(cash_total, 2),
                "received": received or 0,
                "missing": missing,
                "status": status,
                "flagged": flagged,
            },
        })
    return months


BODY_JS = """
function panelFY26(m){
  const row=(l,v,cls)=>`<tr class="${cls||''}"><td>${l}</td><td class="n">${inr(v)}</td><td class="n">${(v/SQFT).toFixed(4)}</td></tr>`;
  const s=m.summary;
  const sum=`<table><tr><th>Item</th><th class="n">Amount (&#8377;)</th><th class="n">Per sqft (&#8377;, ${SQFT.toLocaleString("en-IN")} sqft)</th></tr>`
    +row("Vendor bills",s.bill_expense)+row("Cash withdrawals",s.cash)+row("Total expense (maintenance rate)",s.total_expense,"total")+'</table>';
  const c=m.cash_details;
  const cashCols=[{h:"Date",k:"date"},{h:"Bank narration",k:"description"},{h:"Amount",k:"amount",n:1},{h:"Petty cash details",k:"status"}];
  const msg=c.withdrawals.length?`<p class="${c.flagged?'issues':'okmsg'}">${esc(c.status)}</p>`:'';
  return '<div class="grid">'
    +card("Expense per category (including cash withdrawals)",bars(m.expense_categories))
    +card("Breakup of Others expense",bars(m.others_breakup))
    +card("Summary and maintenance rate",sum)
    +card("Cash transactions",msg+table(cashCols,c.withdrawals,"No cash withdrawals this month"),1)+'</div>';
}
function expenseDebitPanel(){
  const items=DATA.map(m=>({label:m.label,value:m.summary.total_expense}));
  const debitItems=DATA.map(m=>({label:m.label,value:m.debits.net}));
  const rate=v=>(v/SQFT).toFixed(4);
  const rows=DATA.map(m=>({month:m.label,bill:m.summary.bill_expense,cash:m.summary.cash,total:m.summary.total_expense,rate:rate(m.summary.total_expense),gross:m.debits.gross,fd:m.debits.fd,net:m.debits.net,debitRate:rate(m.debits.net)}));
  const sq=`(&#8377;, ${SQFT.toLocaleString("en-IN")} sqft)`;
  const cols=[{h:"Month",k:"month"},{h:"Vendor bills",k:"bill",n:1},{h:"Cash withdrawals",k:"cash",n:1},{h:"Total expense",k:"total",n:1},{h:`Maintenance rate per sqft ${sq}`,k:"rate",r:1},{h:"All bank debits",k:"gross",n:1},{h:"FD transactions (excluded)",k:"fd",n:1},{h:"Net debited",k:"net",n:1},{h:`Debit rate per sqft ${sq}`,k:"debitRate",r:1}];
  return '<div class="grid"><div class="card"><h2>Monthly expense (including cash withdrawals)</h2>'+bars(items,true)+'</div>'
    +'<div class="card"><h2>Money debited from account per month (excluding FD transactions)</h2>'+bars(debitItems,true)+'</div>'
    +'<div class="card wide"><h2>Monthly expense, bank debits and rates</h2>'+table(cols,rows,"No data")+'</div></div>';
}
function cashPanel(){
  const missing=m=>m.cash_details.missing;
  const rows=DATA.map(m=>({month:m.label,count:m.cash_details.withdrawals.length,total:m.cash_details.total,received:m.cash_details.received,missing:missing(m),status:m.cash_details.status,cls:m.cash_details.withdrawals.length?(m.cash_details.flagged?'bad':'ok'):''}));
  const cols=[{h:"Month",k:"month"},{h:"Cash withdrawals",k:"count",r:1},{h:"Withdrawn (bank)",k:"total",n:1},{h:"Petty cash received",k:"received",n:1},{h:"Amount without details",k:"missing",n:1,red:1},{h:"Details",k:"status"}];
  const all=DATA.reduce((a,m)=>a+missing(m),0);
  const note=all?`<p class="issues">Total cash withdrawn without details: <b>&#8377; ${inr(all)}</b></p>`:'<p class="okmsg">Details found for all cash withdrawals</p>';
  return '<div class="grid">'+card("Cash transactions by month (red = no petty cash details found)",note+table(cols,rows,"No data"),1)+'</div>';
}
function renderYear(){
  addTab("Monthly Expense & Debits",expenseDebitPanel());
  addTab("Cash Transactions",cashPanel(),DATA.some(m=>m.cash_details.flagged));
  DATA.forEach(m=>addTab(m.label,panelFY26(m),m.cash_details.flagged));
  show(0);
}
"""


def main():
    data_dir = "test_data"
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--workbook", default=os.path.join(data_dir, "FY26_Expense_Bank_Reconciliation.xlsx"))
    parser.add_argument("--petty", default=os.path.join(data_dir, "2026-2027", "PettyCash.xlsx"))
    parser.add_argument("-o", "--output", default=os.path.join(data_dir, "dashboard_fy26.html"))
    args = parser.parse_args()

    if not os.path.exists(args.workbook):
        raise SystemExit("FY26 workbook not found: %s" % args.workbook)
    if not os.path.exists(args.petty):
        print("Petty cash file not found, all cash withdrawals will show as having no details:", args.petty)
    with open(args.output, "w", encoding="utf-8") as handle:
        handle.write(build_page({LABEL: process_fy26(args.workbook, args.petty)}, BODY_JS))
    print("Wrote", args.output)


if __name__ == "__main__":
    main()
